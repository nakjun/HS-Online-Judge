"""Student roster import and submission persistence."""
import io
import hashlib
import hmac
import re
import secrets
import sqlite3
import time
import zipfile
from contextlib import contextmanager
from pathlib import Path

from openpyxl import load_workbook

DB_PATH = Path(__file__).parent / "judge.db"


@contextmanager
def connect():
    connection = sqlite3.connect(DB_PATH, timeout=10)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    try:
        with connection:
            yield connection
    finally:
        connection.close()


def init_db():
    with connect() as db:
        db.execute("""CREATE TABLE IF NOT EXISTS students (
            student_id TEXT PRIMARY KEY,
            name TEXT NOT NULL,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        )""")
        db.execute("""CREATE TABLE IF NOT EXISTS submissions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_id TEXT NOT NULL REFERENCES students(student_id),
            problem_id TEXT NOT NULL,
            code TEXT NOT NULL,
            verdict TEXT NOT NULL,
            passed INTEGER NOT NULL,
            total INTEGER NOT NULL,
            entry_method TEXT NOT NULL DEFAULT 'UNKNOWN',
            largest_paste_chars INTEGER NOT NULL DEFAULT 0,
            pasted_chars INTEGER NOT NULL DEFAULT 0,
            typed_chars INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        )""")
        db.execute("""CREATE TABLE IF NOT EXISTS student_sessions (
            token_hash TEXT PRIMARY KEY,
            student_id TEXT NOT NULL REFERENCES students(student_id),
            expires_at INTEGER NOT NULL
        )""")
        db.execute("""CREATE TABLE IF NOT EXISTS app_settings (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL,
            updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        )""")
        if db.execute("SELECT 1 FROM app_settings WHERE key='class_code'").fetchone() is None:
            db.execute("INSERT INTO app_settings(key, value) VALUES ('class_code', ?)", (hash_class_code("AS006"),))
            db.execute("DELETE FROM student_sessions")
        existing = {row["name"] for row in db.execute("PRAGMA table_info(submissions)")}
        for column, definition in {
            "entry_method": "TEXT NOT NULL DEFAULT 'UNKNOWN'",
            "largest_paste_chars": "INTEGER NOT NULL DEFAULT 0",
            "pasted_chars": "INTEGER NOT NULL DEFAULT 0",
            "typed_chars": "INTEGER NOT NULL DEFAULT 0",
        }.items():
            if column not in existing:
                db.execute(f"ALTER TABLE submissions ADD COLUMN {column} {definition}")


def cell_text(cell):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        raise ValueError("학번과 이름에는 참/거짓 값을 사용할 수 없습니다.")
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError("학번은 정수 또는 텍스트여야 합니다.")
        value = int(value)
    if isinstance(value, int) and cell.number_format and set(cell.number_format) == {"0"}:
        return str(value).zfill(len(cell.number_format))
    return str(value).strip()


def parse_roster(content):
    if len(content) > 2_000_000:
        raise ValueError("엑셀 파일은 2MB 이하여야 합니다.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 20_000_000:
                raise ValueError("압축을 푼 엑셀 데이터가 너무 큽니다.")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, OSError, ValueError) as exc:
        raise ValueError(f"올바른 .xlsx 파일이 아닙니다: {exc}") from exc
    try:
        sheet = workbook.active
        rows = sheet.iter_rows()
        header = next(rows, None)
        if header is None:
            raise ValueError("엑셀 파일이 비어 있습니다.")
        names = [cell_text(cell) for cell in header]
        if "학번" not in names or "이름" not in names:
            raise ValueError("첫 행에 '학번'과 '이름' 열이 필요합니다.")
        id_col, name_col = names.index("학번"), names.index("이름")
        students = {}
        for row_number, row in enumerate(rows, start=2):
            if row_number > 10_001:
                raise ValueError("학생은 최대 10,000명까지 가져올 수 있습니다.")
            student_id = cell_text(row[id_col]) if id_col < len(row) else ""
            name = cell_text(row[name_col]) if name_col < len(row) else ""
            if not student_id and not name:
                continue
            if not student_id or not name or len(student_id) > 40 or len(name) > 100:
                raise ValueError(f"{row_number}행의 학번 또는 이름이 비어 있거나 너무 깁니다.")
            if student_id in students and students[student_id] != name:
                raise ValueError(f"{row_number}행: 같은 학번에 다른 이름이 있습니다.")
            students[student_id] = name
        if not students:
            raise ValueError("가져올 학생이 없습니다.")
        return students
    finally:
        workbook.close()


def import_students(content):
    students = parse_roster(content)
    with connect() as db:
        existing = {row["student_id"]: row["name"] for row in db.execute("SELECT student_id, name FROM students")}
        for student_id, name in students.items():
            db.execute("INSERT INTO students(student_id, name) VALUES (?, ?) ON CONFLICT(student_id) DO UPDATE SET name=excluded.name", (student_id, name))
    return {"total": len(students), "created": sum(sid not in existing for sid in students),
            "updated": sum(sid in existing and existing[sid] != name for sid, name in students.items())}


def student_exists(student_id, name):
    with connect() as db:
        return db.execute("SELECT 1 FROM students WHERE student_id=? AND name=?", (student_id, name)).fetchone() is not None


def hash_class_code(code):
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", code.encode("utf-8"), salt, 200_000)
    return salt.hex() + ":" + digest.hex()


def verify_class_code(code):
    if not isinstance(code, str):
        return False
    with connect() as db:
        row = db.execute("SELECT value FROM app_settings WHERE key='class_code'").fetchone()
    if not row:
        return False
    salt_hex, digest_hex = row["value"].split(":", 1)
    actual = hashlib.pbkdf2_hmac("sha256", code.encode("utf-8"), bytes.fromhex(salt_hex), 200_000)
    return hmac.compare_digest(actual, bytes.fromhex(digest_hex))


def class_code_status():
    with connect() as db:
        row = db.execute("SELECT updated_at FROM app_settings WHERE key='class_code'").fetchone()
    return {"is_default": verify_class_code("AS006"), "updated_at": row["updated_at"] if row else None}


def set_class_code(code):
    if not isinstance(code, str) or not re.fullmatch(r"[A-Za-z0-9_-]{4,32}", code):
        raise ValueError("수업 코드는 영문, 숫자, _ 또는 -를 사용해 4~32자로 입력하세요.")
    with connect() as db:
        db.execute("UPDATE app_settings SET value=?, updated_at=CURRENT_TIMESTAMP WHERE key='class_code'",
                   (hash_class_code(code),))
        db.execute("DELETE FROM student_sessions")
    return class_code_status()


def create_student_session(student_id):
    token = secrets.token_urlsafe(32)
    token_hash = hashlib.sha256(token.encode("ascii")).hexdigest()
    with connect() as db:
        db.execute("DELETE FROM student_sessions WHERE expires_at < ?", (int(time.time()),))
        db.execute("INSERT INTO student_sessions(token_hash, student_id, expires_at) VALUES (?, ?, ?)",
                   (token_hash, student_id, int(time.time()) + 30 * 24 * 60 * 60))
    return token


def session_student(token):
    if not isinstance(token, str) or not 20 <= len(token) <= 100:
        return None
    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
    with connect() as db:
        row = db.execute("""SELECT st.student_id, st.name FROM student_sessions ss
            JOIN students st ON st.student_id=ss.student_id
            WHERE ss.token_hash=? AND ss.expires_at>?""", (token_hash, int(time.time()))).fetchone()
        return dict(row) if row else None


def delete_student_session(token):
    if not token:
        return
    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
    with connect() as db:
        db.execute("DELETE FROM student_sessions WHERE token_hash=?", (token_hash,))


def student_progress(student_id, weeks, problems):
    with connect() as db:
        state = {}
        for row in db.execute("SELECT problem_id, verdict FROM submissions WHERE student_id=? ORDER BY id", (student_id,)):
            item = state.setdefault(row["problem_id"], {"attempts": 0, "solved": False})
            item["attempts"] += 1
            item["solved"] |= row["verdict"] == "ACCEPTED"
    return [{"id": week["id"], "title": week["title"], "problems": [
        {"id": problem["id"], "title": problem["title"],
         "attempts": state.get(problem["id"], {}).get("attempts", 0),
         "solved": state.get(problem["id"], {}).get("solved", False)}
        for problem in problems if (problem.get("week_id") or "week-1") == week["id"]]}
        for week in weeks]


def classify_entry(code, entry):
    if not isinstance(entry, dict):
        return "UNKNOWN", 0, 0, 0
    numbers = []
    for field in ("largest_paste_chars", "pasted_chars", "typed_chars"):
        value = entry.get(field, 0)
        if not isinstance(value, int) or isinstance(value, bool) or not 0 <= value <= 1_000_000:
            return "UNKNOWN", 0, 0, 0
        numbers.append(value)
    largest, pasted, typed = numbers
    if largest >= 20 and largest >= len(code) * 0.9:
        method = "FULL_PASTE"
    elif entry.get("file_upload") is True:
        method = "FILE_UPLOAD"
    elif pasted > 0:
        method = "PARTIAL_PASTE"
    elif typed > 0:
        method = "TYPED"
    else:
        method = "UNKNOWN"
    return method, largest, pasted, typed


def save_submission(student_id, problem_id, code, verdict, passed, total, entry=None):
    method, largest, pasted, typed = classify_entry(code, entry)
    with connect() as db:
        cursor = db.execute("""INSERT INTO submissions(student_id, problem_id, code, verdict, passed, total,
            entry_method, largest_paste_chars, pasted_chars, typed_chars) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                            (student_id, problem_id, code, verdict, passed, total, method, largest, pasted, typed))
        return cursor.lastrowid


def list_students():
    with connect() as db:
        return [dict(row) for row in db.execute("SELECT student_id, name FROM students ORDER BY student_id")]


def list_submissions():
    with connect() as db:
        return [dict(row) for row in db.execute("""SELECT s.id, s.student_id, st.name, s.problem_id,
            s.verdict, s.passed, s.total, s.entry_method, s.largest_paste_chars, s.created_at
            FROM submissions s JOIN students st ON st.student_id=s.student_id
            ORDER BY s.id DESC LIMIT 200""")]


def get_submission(submission_id):
    with connect() as db:
        row = db.execute("""SELECT s.*, st.name FROM submissions s
            JOIN students st ON st.student_id=s.student_id WHERE s.id=?""", (submission_id,)).fetchone()
        return dict(row) if row else None


def progress_page(weeks, problems, query="", offset=0, limit=100):
    """Return one roster page with per-week, per-problem submission state."""
    with connect() as db:
        pattern = "%" + query.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%"
        where = "WHERE student_id LIKE ? ESCAPE '\\' OR name LIKE ? ESCAPE '\\'"
        total_students = db.execute(f"SELECT COUNT(*) FROM students {where}", (pattern, pattern)).fetchone()[0]
        students = [dict(row) for row in db.execute(
            f"SELECT student_id, name FROM students {where} ORDER BY student_id LIMIT ? OFFSET ?",
            (pattern, pattern, limit, offset))]
        ids = [student["student_id"] for student in students]
        submission_state = {}
        if ids:
            placeholders = ",".join("?" for _ in ids)
            records = db.execute(
                f"SELECT student_id, problem_id, verdict FROM submissions WHERE student_id IN ({placeholders}) ORDER BY id",
                ids)
            for record in records:
                key = (record["student_id"], record["problem_id"])
                state = submission_state.setdefault(key, {"attempts": 0, "accepted": False, "latest_verdict": ""})
                state["attempts"] += 1
                state["accepted"] |= record["verdict"] == "ACCEPTED"
                state["latest_verdict"] = record["verdict"]

    problems_by_week = {week["id"]: [] for week in weeks}
    for problem in problems:
        problems_by_week.setdefault(problem.get("week_id") or "week-1", []).append(problem)
    rows = []
    for student in students:
        cells = []
        for week in weeks:
            details = []
            for problem in problems_by_week.get(week["id"], []):
                state = submission_state.get((student["student_id"], problem["id"]))
                details.append({"problem_id": problem["id"], "title": problem["title"],
                                "attempts": state["attempts"] if state else 0,
                                "accepted": state["accepted"] if state else False,
                                "latest_verdict": state["latest_verdict"] if state else ""})
            cells.append({"week_id": week["id"], "total": len(details),
                          "submitted": sum(item["attempts"] > 0 for item in details),
                          "accepted": sum(item["accepted"] for item in details), "problems": details})
        rows.append({**student, "weeks": cells})
    return {"weeks": weeks, "total_students": total_students, "offset": offset, "limit": limit, "rows": rows}
